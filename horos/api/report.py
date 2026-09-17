"""Training-run reports (E8: chart export).

`build_training_report` gathers everything a run recorded — config, derived
hyperparameters with reasons, the metric series from events.jsonl, the best
checkpoint's scores, the verdict, evaluation reports, the dataset snapshot's
composition — into one pydantic object. The renderers turn that object into:

  png   one 1920x1080 (16:9) dashboard image that drops straight into a slide
  pdf   the dashboard, then hyperparameter / evaluation / findings pages
  xlsx  one workbook: Summary, Metrics (per epoch), Hyperparameters, Classes,
        Evaluation, Verdict

matplotlib and openpyxl are imported lazily inside the renderers: they ship
with the ML stack (`horos install`), never with the annotation-only core.
"""

from __future__ import annotations

import json
import re
import textwrap
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

import horos
from horos.core.project import Project
from horos.errors import ProjectError

ReportFormat = Literal["png", "pdf", "xlsx"]
REPORT_FORMATS: tuple[str, ...] = ("png", "pdf", "xlsx")

#: the metric keys a reader acts on (same filter as the Train page's chips)
_SCORE_RE = re.compile(r"(mAP|mAR|F1|precision|recall)", re.IGNORECASE)
_MAP_RE = re.compile(r"map", re.IGNORECASE)

PNG_SIZE = (1920, 1080)  # 16:9
_FIGSIZE = (16, 9)
_DPI = 120


class MetricSeries(BaseModel):
    key: str
    points: list[tuple[int, float]]  # (epoch, value)


class ReportHparam(BaseModel):
    name: str
    value: Any
    reason: str = ""
    overridden: bool = False


class ReportClassEval(BaseModel):
    name: str
    instances: int
    ap: float
    ap50: float


class ReportEval(BaseModel):
    split: str
    num_images: int
    num_instances: int
    map_5095: float
    map_50: float
    map_75: float
    mar_100: float
    per_class: list[ReportClassEval] = Field(default_factory=list)


class ReportFinding(BaseModel):
    severity: str
    title: str
    detail: str
    suggestion: str


class TrainingReport(BaseModel):
    run_id: str
    model: str
    model_display: str
    license: str
    state: str
    created_at: str
    generated_at: str
    epochs_completed: int | None = None
    epochs_planned: int | None = None
    categories: list[str] = Field(default_factory=list)
    include_background: bool = False
    seed: int | None = None
    checkpoint_criterion: str = ""
    device: str | None = None
    dataset_images: int = 0
    dataset_splits: dict[str, int] = Field(default_factory=dict)
    #: train-split instances per class, from the run's dataset snapshot
    class_instances: dict[str, int] = Field(default_factory=dict)
    hparams: list[ReportHparam] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)
    #: (step, value) per metric, in the backend's own epoch numbering — rfdetr
    #: reports 0-based epochs; the charts plot the raw steps like the Train page
    series: list[MetricSeries] = Field(default_factory=list)
    best_epoch: int | None = None  # 0-based, as the backend reports it
    #: scores at the best checkpoint (or the last epoch when no best is known)
    final_metrics: dict[str, float] = Field(default_factory=dict)
    verdict_summary: str = ""
    findings: list[ReportFinding] = Field(default_factory=list)
    evals: list[ReportEval] = Field(default_factory=list)
    checkpoint: str | None = None
    error: str | None = None

    def series_for(self, key: str) -> list[tuple[int, float]] | None:
        return next((s.points for s in self.series if s.key == key), None)

    @property
    def epoch_offset(self) -> int:
        """Backends that count epochs from 0 (rfdetr) are shown 1-based, like
        the Train page's "best checkpoint (epoch N)" note; 1-based streams
        (the fakes) are shown as they are."""
        steps = [e for s in self.series for e, _ in s.points]
        return 1 if steps and min(steps) == 0 else 0

    def display_epoch(self, step: int) -> int:
        return step + self.epoch_offset


# ------------------------------------------------------------------ gathering


def _snapshot_class_instances(run_dir: Path) -> dict[str, int]:
    gt_path = run_dir / "dataset" / "train" / "_annotations.coco.json"
    if not gt_path.is_file():
        return {}
    try:
        gt = json.loads(gt_path.read_text("utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    names = {c["id"]: c["name"] for c in gt.get("categories", [])}
    counts = {name: 0 for name in names.values()}
    for ann in gt.get("annotations", []):
        name = names.get(ann.get("category_id"))
        if name is not None:
            counts[name] += 1
    return counts


def _series_from_events(events: list[dict]) -> dict[str, list[tuple[int, float]]]:
    series: dict[str, list[tuple[int, float]]] = {}
    for event in events:
        if event.get("type") != "metrics":
            continue
        step = event.get("step")
        for key, value in (event.get("metrics") or {}).items():
            if isinstance(value, int | float) and isinstance(step, int):
                series.setdefault(key, []).append((step, float(value)))
    return series


def run_scores(
    raw_series: dict[str, list[tuple[int, float]]],
) -> tuple[int | None, dict[str, float]]:
    """(best epoch, scores at the best checkpoint) from a run's metric series.

    Scores are the mAP/mAR/F1/precision/recall keys taken at the epoch the
    backend reported as best (or the last epoch when no best is known), plus
    the final loss values. Shared by the report and `horos models`."""
    best_pts = raw_series.get("best/epoch")
    best_epoch = int(best_pts[-1][1]) if best_pts else None
    final: dict[str, float] = {}
    for key, points in raw_series.items():
        if key.startswith("best/") or key.startswith("__"):  # backend-internal bookkeeping
            continue
        if _SCORE_RE.search(key) and "loss" not in key.lower():
            at_best = next((v for e, v in points if e == best_epoch), None)
            final[key] = at_best if at_best is not None else points[-1][1]
    for key in ("train/loss", "val/loss", "loss"):
        if key in raw_series:
            final[key] = raw_series[key][-1][1]
    return best_epoch, final


def build_training_report(project: Project, run_id: str) -> TrainingReport:
    """Everything the renderers need, gathered once from the run directory."""
    from horos.api.train import _read_events, _run_dir, _split_counts, read_record
    from horos.api.verdict import build_verdict
    from horos.core.registry import get_model_info
    from horos.errors import UnknownModelError

    run_dir = _run_dir(project, run_id)
    record = read_record(run_dir)
    events, _ = _read_events(run_dir)
    splits = _split_counts(run_dir, record)

    try:
        info = get_model_info(record.model)
        display, license_ = info.display_name, info.weights_license
    except UnknownModelError:
        display, license_ = record.model, "unknown"

    raw_series = _series_from_events(events)
    best_epoch, final = run_scores(raw_series)

    verdict = build_verdict(record, events, splits)

    evals: list[ReportEval] = []
    for split in ("valid", "test", "train"):
        path = run_dir / "eval" / f"{split}.json"
        if not path.is_file():
            continue
        try:
            data = json.loads(path.read_text("utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        evals.append(
            ReportEval(
                split=split,
                num_images=data.get("num_images", 0),
                num_instances=data.get("num_instances", 0),
                map_5095=data.get("map_5095", 0.0),
                map_50=data.get("map_50", 0.0),
                map_75=data.get("map_75", 0.0),
                mar_100=data.get("mar_100", 0.0),
                per_class=[
                    ReportClassEval(
                        name=c["name"], instances=c.get("instances", 0),
                        ap=c.get("ap", 0.0), ap50=c.get("ap50", 0.0),
                    )
                    for c in data.get("per_class", [])
                ],
            )
        )

    config = record.config or {}
    return TrainingReport(
        run_id=record.run_id,
        model=record.model,
        model_display=display,
        license=license_,
        state=record.state,
        created_at=record.created_at,
        generated_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        epochs_completed=record.epochs_completed,
        epochs_planned=config.get("epochs"),
        categories=list(record.dataset_classes or config.get("categories") or []),
        include_background=bool(config.get("include_background", False)),
        seed=config.get("seed"),
        checkpoint_criterion=str(config.get("checkpoint_criterion", "")),
        device=record.device,
        dataset_images=record.dataset_images,
        dataset_splits=dict(splits),
        class_instances=_snapshot_class_instances(run_dir),
        hparams=[
            ReportHparam(
                name=h.name, value=h.value, reason=h.reason, overridden=h.overridden
            )
            for h in record.hparams
        ],
        notes=list(record.hparam_notes),
        series=[
            MetricSeries(key=k, points=v)
            for k, v in sorted(raw_series.items())
            if not k.startswith("__")
        ],
        best_epoch=best_epoch,
        final_metrics=final,
        verdict_summary=verdict.summary,
        findings=[
            ReportFinding(
                severity=str(f.severity), title=f.title, detail=f.detail,
                suggestion=f.suggestion,
            )
            for f in verdict.findings
        ],
        evals=evals,
        checkpoint=record.checkpoint,
        error=record.error,
    )


# ------------------------------------------------------------------ rendering


def _fmt(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.4g}" if abs(value) < 1000 else f"{value:,.0f}"
    if isinstance(value, dict | list):
        text = json.dumps(value)
        return text if len(text) <= 40 else text[:37] + "…"
    return str(value)


def _wrap(text: str, width: int) -> str:
    return "\n".join(textwrap.wrap(text, width)) if text else ""


def _import_matplotlib():
    try:
        import matplotlib

        matplotlib.use("Agg", force=False)
        from matplotlib.figure import Figure
        from matplotlib.gridspec import GridSpec
    except ImportError as exc:  # pragma: no cover — depends on the environment
        raise ProjectError(
            "Report export needs matplotlib, which ships with the ML stack. "
            "Run 'horos install' (or pip install matplotlib) and retry."
        ) from exc
    return Figure, GridSpec


_COLORS = ["#4dabf7", "#f59f00", "#3cb44b", "#e6194b", "#911eb4", "#42d4f4", "#f032e6"]
_TEXT = "#111111"
_DIM = "#666666"


def _style(ax) -> None:
    from matplotlib.ticker import MaxNLocator

    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.grid(True, alpha=0.25)
    ax.tick_params(labelsize=9)
    ax.xaxis.set_major_locator(MaxNLocator(integer=True))  # epochs are whole numbers


def _dashboard_figure(report: TrainingReport):
    """The 16:9 dashboard: title, loss curves, mAP curves, scores, dataset
    composition, key hyperparameters, conclusion."""
    Figure, GridSpec = _import_matplotlib()
    fig = Figure(figsize=_FIGSIZE, dpi=_DPI)
    fig.patch.set_facecolor("white")
    gs = GridSpec(
        4, 3, figure=fig, height_ratios=[0.7, 1.55, 1.55, 1.35],
        left=0.05, right=0.97, top=0.96, bottom=0.06, hspace=0.55, wspace=0.28,
    )

    # -- title row
    title_ax = fig.add_subplot(gs[0, :])
    title_ax.axis("off")
    title_ax.text(
        0, 0.85, f"{report.model_display} · training run {report.run_id}",
        fontsize=20, fontweight="bold", color=_TEXT, va="top",
    )
    epochs = (
        f"{report.epochs_completed}/{report.epochs_planned} epochs"
        if report.epochs_completed is not None and report.epochs_planned
        else f"{report.epochs_planned or '?'} epochs planned"
    )
    classes = ", ".join(report.categories[:6]) + (
        f" +{len(report.categories) - 6}" if len(report.categories) > 6 else ""
    )
    splits = " · ".join(f"{k} {v}" for k, v in report.dataset_splits.items() if v)
    subtitle = (
        f"state: {report.state}   ·   {epochs}   ·   {report.dataset_images} images "
        f"({splits})   ·   license: {report.license}   ·   created {report.created_at[:19]}\n"
        f"{len(report.categories)} classes: {classes}"
    )
    title_ax.text(0, 0.28, subtitle, fontsize=10.5, color=_DIM, va="top", linespacing=1.4)

    # -- loss curves
    ax_loss = fig.add_subplot(gs[1:3, 0:2])
    loss_keys = [k for k in ("train/loss", "val/loss") if report.series_for(k)]
    if not loss_keys:
        loss_keys = [s.key for s in report.series if s.key.lower() == "loss"]
    shown_best = None if report.best_epoch is None else report.display_epoch(report.best_epoch)
    for i, key in enumerate(loss_keys):
        pts = report.series_for(key) or []
        ax_loss.plot([report.display_epoch(e) for e, _ in pts], [v for _, v in pts],
                     color=_COLORS[i], linewidth=2, label=key)
    if shown_best is not None:
        ax_loss.axvline(shown_best, color="#3cb44b", linestyle="--",
                        linewidth=1.2, label=f"best checkpoint (epoch {shown_best})")
    ax_loss.set_title("Loss", fontsize=12, loc="left", color=_TEXT)
    ax_loss.set_xlabel("epoch", fontsize=9)
    if loss_keys or report.best_epoch is not None:
        ax_loss.legend(fontsize=9, frameon=False)
    else:
        ax_loss.text(0.5, 0.5, "no loss series recorded", ha="center", va="center",
                     color=_DIM, transform=ax_loss.transAxes)
    _style(ax_loss)

    # -- mAP curves
    ax_map = fig.add_subplot(gs[1, 2])
    map_keys = [s.key for s in report.series
                if _MAP_RE.search(s.key) and "loss" not in s.key.lower()
                and not s.key.startswith("best/")][:4]
    for i, key in enumerate(map_keys):
        pts = report.series_for(key) or []
        ax_map.plot([report.display_epoch(e) for e, _ in pts], [v for _, v in pts],
                    color=_COLORS[(i + 2) % 7], linewidth=1.8, label=key)
    if shown_best is not None and map_keys:
        ax_map.axvline(shown_best, color="#3cb44b", linestyle="--", linewidth=1)
    ax_map.set_title("Validation mAP", fontsize=12, loc="left", color=_TEXT)
    if map_keys:
        ax_map.legend(fontsize=8, frameon=False)
    else:
        ax_map.text(0.5, 0.5, "no mAP series recorded", ha="center", va="center",
                    color=_DIM, transform=ax_map.transAxes)
    _style(ax_map)

    # -- scores at the best checkpoint
    ax_scores = fig.add_subplot(gs[2, 2])
    ax_scores.axis("off")
    where = (
        f"scores at the best checkpoint (epoch {shown_best})"
        if shown_best is not None else "scores at the last epoch"
    )
    ax_scores.set_title(where, fontsize=11, loc="left", color=_TEXT)
    rows = sorted(report.final_metrics.items(), key=lambda kv: ("loss" in kv[0], kv[0]))[:9]
    if rows:
        for i, (key, value) in enumerate(rows):
            y = 0.92 - i * 0.105
            ax_scores.text(0.0, y, key, fontsize=9.5, color=_DIM, va="top",
                           transform=ax_scores.transAxes)
            ax_scores.text(1.0, y, _fmt(value), fontsize=10.5, color=_TEXT, va="top",
                           ha="right", fontweight="bold", transform=ax_scores.transAxes)
    else:
        ax_scores.text(0, 0.9, "no metrics recorded", fontsize=9.5, color=_DIM, va="top",
                       transform=ax_scores.transAxes)

    # -- dataset composition
    ax_cls = fig.add_subplot(gs[3, 0])
    counts = sorted(report.class_instances.items(), key=lambda kv: -kv[1])[:10]
    if counts:
        names = [n for n, _ in counts][::-1]
        values = [v for _, v in counts][::-1]
        # long class names need room on the left: shrink the axes by the label width
        longest = max(len(n) for n in names)
        shift = min(0.1, 0.0026 * longest)
        x0, y0, w, h = ax_cls.get_position().bounds
        ax_cls.set_position([x0 + shift, y0, w - shift - 0.02, h])
        ax_cls.barh(names, values, color=_COLORS[0])
        for y, v in enumerate(values):
            ax_cls.text(v, y, f" {v}", va="center", fontsize=8.5, color=_DIM)
        ax_cls.tick_params(labelsize=8.5)
        ax_cls.set_xticks([])
        ax_cls.set_xlim(0, max(values) * 1.18)
    else:
        ax_cls.text(0.5, 0.5, "no snapshot", ha="center", va="center", color=_DIM,
                    transform=ax_cls.transAxes)
    ax_cls.set_title("Training instances per class", fontsize=11, loc="left", color=_TEXT)
    for side in ("top", "right", "bottom"):
        ax_cls.spines[side].set_visible(False)

    # -- key hyperparameters
    ax_hp = fig.add_subplot(gs[3, 1])
    ax_hp.axis("off")
    ax_hp.set_title("Hyperparameters", fontsize=11, loc="left", color=_TEXT)
    preferred = ["epochs", "batch_size", "grad_accum_steps", "resolution", "lr",
                 "warmup_epochs", "early_stopping_patience", "augmentation_backend"]
    by_name = {h.name: h for h in report.hparams}
    shown = [by_name[n] for n in preferred if n in by_name]
    shown += [h for h in report.hparams if h.name not in preferred][: max(0, 9 - len(shown))]
    for i, h in enumerate(shown[:9]):
        y = 0.95 - i * 0.11
        ax_hp.text(0.0, y, h.name, fontsize=9, color=_DIM, va="top", transform=ax_hp.transAxes)
        ax_hp.text(1.0, y, _fmt(h.value) + (" ✎" if h.overridden else ""), fontsize=9.5,
                   color=_TEXT, va="top", ha="right", transform=ax_hp.transAxes)
    if not shown:
        ax_hp.text(0, 0.9, "none recorded", fontsize=9.5, color=_DIM, va="top",
                   transform=ax_hp.transAxes)

    # -- conclusion
    ax_v = fig.add_subplot(gs[3, 2])
    ax_v.axis("off")
    ax_v.set_title("Conclusion", fontsize=11, loc="left", color=_TEXT)
    line_h = 0.085  # axes fraction per text line at these font sizes
    y = 0.95
    summary_lines = textwrap.wrap(report.verdict_summary or report.error or "—", 46) or ["—"]
    ax_v.text(0, y, "\n".join(summary_lines), fontsize=9.5, color=_TEXT, va="top",
              fontweight="bold", transform=ax_v.transAxes)
    y -= line_h * len(summary_lines) + 0.05
    for finding in report.findings[:3]:
        lines = textwrap.wrap(f"• [{finding.severity}] {finding.title}", 48) or ["•"]
        if y - line_h * len(lines) < 0.0:
            break
        ax_v.text(0, y, "\n".join(lines), fontsize=8.8, color=_DIM, va="top",
                  transform=ax_v.transAxes)
        y -= line_h * len(lines) + 0.03

    fig.text(0.97, 0.015, f"generated by horos · {report.generated_at}",
             fontsize=8, color=_DIM, ha="right")
    return fig


def render_png(report: TrainingReport, path: Path) -> Path:
    """One 1920x1080 image (16:9), ready for a slide or a document."""
    fig = _dashboard_figure(report)
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=_DPI, facecolor="white")
    return path


#: text lines that fit under a page title at the table font size
_LINES_PER_PAGE = 40
#: characters that fit across the full table width at the table font size
_CHARS_PER_ROW = 210


def _table_pages(Figure, title: str, columns: list[str], rows: list[list[str]],
                 col_widths: list[float], note: str = ""):
    """Yield one figure per page. Cell text is wrapped to its column width and
    every row is as tall as its longest cell — matplotlib's table gives all
    rows the same height, which is what made long reasons spill over."""
    line_h = 1.0 / (_LINES_PER_PAGE + 4)  # axes fraction per text line
    note_lines = textwrap.wrap(note, 160) if note else []

    def wrap_row(row: list[str]) -> tuple[list[str], int]:
        cells = []
        for text, width in zip(row, col_widths, strict=True):
            lines = textwrap.wrap(str(text), max(8, int(width * _CHARS_PER_ROW) - 2)) or [""]
            cells.append("\n".join(lines))
        return cells, max(c.count("\n") + 1 for c in cells)

    wrapped = [wrap_row(r) for r in rows]
    budget = _LINES_PER_PAGE - 2 * len(note_lines) - 3  # header + breathing room
    pages: list[list[tuple[list[str], int]]] = [[]]
    used = 0
    for cells, lines in wrapped:
        if pages[-1] and used + lines + 1 > budget:
            pages.append([])
            used = 0
            budget = _LINES_PER_PAGE - 3
        pages[-1].append((cells, lines))
        used += lines + 1

    for index, page in enumerate(pages):
        fig = Figure(figsize=_FIGSIZE, dpi=_DPI)
        fig.patch.set_facecolor("white")
        ax = fig.add_subplot(111)
        ax.axis("off")
        suffix = f"  ({index + 1}/{len(pages)})" if len(pages) > 1 else ""
        fig.text(0.05, 0.94, title + suffix, fontsize=18, fontweight="bold", color=_TEXT, va="top")
        top = 0.86
        if index == 0 and note_lines:
            fig.text(0.05, 0.885, "\n".join(note_lines), fontsize=9.5, color=_DIM, va="top",
                     linespacing=1.4)
            top -= 0.028 * len(note_lines)
        if not page:
            fig.text(0.05, top - 0.05, "nothing recorded", fontsize=11, color=_DIM)
            yield fig
            continue
        table = ax.table(
            cellText=[cells for cells, _ in page], colLabels=columns, loc="upper left",
            cellLoc="left", colLoc="left", colWidths=col_widths,
        )
        table.auto_set_font_size(False)
        table.set_fontsize(8.5)
        heights = [2] + [lines + 1 for _, lines in page]  # header, then each row
        # the axes is shrunk to the table's total height (figure fraction), so
        # each cell takes its share of the axes — never line_h again on top
        total = sum(heights) * line_h
        ax.set_position([0.05, max(0.02, top - total), 0.9, total])
        for (row, _col), cell in table.get_celld().items():
            cell.set_edgecolor("#dddddd")
            cell.set_height(heights[row] / sum(heights))
            cell.PAD = 0.02
            if row == 0:
                cell.set_facecolor("#f1f3f5")
                cell.set_text_props(fontweight="bold", color=_TEXT)
        yield fig


def render_pdf(report: TrainingReport, path: Path) -> Path:
    """Dashboard first, then the detail pages."""
    Figure, _ = _import_matplotlib()
    from matplotlib.backends.backend_pdf import PdfPages

    path.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(path) as pdf:
        pdf.savefig(_dashboard_figure(report))

        rows = [
            [h.name, _fmt(h.value), "override" if h.overridden else "derived", h.reason]
            for h in report.hparams
        ]
        for fig in _table_pages(
            Figure, "Hyperparameters and why they were chosen",
            ["name", "value", "origin", "reason"], rows, [0.16, 0.12, 0.09, 0.63],
            note=" | ".join(report.notes) if report.notes else "",
        ):
            pdf.savefig(fig)

        eval_rows: list[list[str]] = []
        for ev in report.evals:
            eval_rows.append([ev.split, "(all)", str(ev.num_instances),
                              f"{ev.map_50:.4f}", f"{ev.map_5095:.4f}",
                              f"mAP75 {ev.map_75:.4f} · mAR100 {ev.mar_100:.4f} · "
                              f"{ev.num_images} images"])
            for c in ev.per_class:
                eval_rows.append([ev.split, c.name, str(c.instances),
                                  f"{c.ap50:.4f}", f"{c.ap:.4f}", ""])
        for fig in _table_pages(
            Figure, "Evaluation on the held-out splits",
            ["split", "class", "instances", "AP50", "AP@[.50:.95]", "notes"],
            eval_rows, [0.08, 0.2, 0.1, 0.1, 0.12, 0.4],
            note="" if report.evals else
            "No evaluation report yet — run the Evaluate page (or 'horos evaluate') "
            "on this run to add per-class numbers.",
        ):
            pdf.savefig(fig)

        finding_rows = [[f.severity, f.title, f.detail, f.suggestion] for f in report.findings]
        for fig in _table_pages(
            Figure, "Conclusion and suggestions",
            ["severity", "finding", "detail", "suggestion"], finding_rows,
            [0.09, 0.24, 0.37, 0.3], note=report.verdict_summary,
        ):
            pdf.savefig(fig)
        info = pdf.infodict()
        info["Title"] = f"horos training report {report.run_id}"
        info["Creator"] = "horos"
    return path


def render_xlsx(report: TrainingReport, path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover — depends on the environment
        raise ProjectError(
            "Excel export needs openpyxl, which ships with the ML stack. "
            "Run 'horos install' (or pip install openpyxl) and retry."
        ) from exc

    bold = Font(bold=True)

    def header(ws, names: list[str]) -> None:
        ws.append(names)
        for cell in ws[1]:
            cell.font = bold

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    header(ws, ["field", "value"])
    for key, value in [
        ("run_id", report.run_id), ("model", report.model),
        ("model_display", report.model_display), ("license", report.license),
        ("state", report.state), ("created_at", report.created_at),
        ("generated_at", report.generated_at),
        ("epochs_completed", report.epochs_completed),
        ("epochs_planned", report.epochs_planned),
        ("classes", ", ".join(report.categories)),
        ("include_background", report.include_background),
        ("seed", report.seed), ("checkpoint_criterion", report.checkpoint_criterion),
        ("device", report.device), ("dataset_images", report.dataset_images),
        *[(f"split_{k}", v) for k, v in report.dataset_splits.items()],
        ("best_epoch", None if report.best_epoch is None
         else report.display_epoch(report.best_epoch)),
        *[(f"score:{k}", v) for k, v in sorted(report.final_metrics.items())],
        ("checkpoint", report.checkpoint), ("verdict", report.verdict_summary),
    ]:
        ws.append([key, value])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70

    ws = wb.create_sheet("Metrics")
    keys = [s.key for s in report.series]
    header(ws, ["epoch", *keys])
    epochs = sorted({e for s in report.series for e, _ in s.points})
    lookup = {s.key: dict(s.points) for s in report.series}
    for epoch in epochs:
        ws.append([report.display_epoch(epoch), *[lookup[k].get(epoch) for k in keys]])

    ws = wb.create_sheet("Hyperparameters")
    header(ws, ["name", "value", "origin", "reason"])
    for h in report.hparams:
        ws.append([h.name, h.value if isinstance(h.value, int | float | str | bool)
                   else json.dumps(h.value),
                   "override" if h.overridden else "derived", h.reason])
    for note in report.notes:
        ws.append(["note", "", "", note])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["D"].width = 100

    ws = wb.create_sheet("Classes")
    header(ws, ["class", "train_instances"])
    for name, count in sorted(report.class_instances.items(), key=lambda kv: -kv[1]):
        ws.append([name, count])

    ws = wb.create_sheet("Evaluation")
    header(ws, ["split", "class", "instances", "AP50", "AP", "mAP75", "mAR100", "images"])
    for ev in report.evals:
        ws.append([ev.split, "(all)", ev.num_instances, ev.map_50, ev.map_5095,
                   ev.map_75, ev.mar_100, ev.num_images])
        for c in ev.per_class:
            ws.append([ev.split, c.name, c.instances, c.ap50, c.ap, None, None, None])

    ws = wb.create_sheet("Verdict")
    header(ws, ["severity", "title", "detail", "suggestion"])
    ws.append(["summary", report.verdict_summary, "", ""])
    for f in report.findings:
        ws.append([f.severity, f.title, f.detail, f.suggestion])
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 60

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


RENDERERS = {"png": render_png, "pdf": render_pdf, "xlsx": render_xlsx}


def render_report(report: TrainingReport, format: str, path: Path) -> Path:
    if format not in RENDERERS:
        raise ProjectError(
            f"Unsupported report format '{format}' ({'|'.join(REPORT_FORMATS)})"
        )
    return RENDERERS[format](report, path)


# ------------------------------------------------- evaluation chart (E6-T14)

#: sequential blue, light -> dark: the matrix shades magnitude, so it is ONE
#: hue by lightness, never a rainbow. Cells carry their count as text, so the
#: shade is a scan aid and never the only way to read a number.
_SEQ = ["#cde2fb", "#b7d3f6", "#9ec5f4", "#86b6ef", "#6da7ec", "#5598e7",
        "#3987e5", "#2a78d6", "#256abf", "#1c5cab", "#184f95", "#104281", "#0d366b"]
#: the diagonal is "correct" — outlined, not recolored, so hue stays magnitude
_DIAG = "#0ca30c"
_RULE = "#d9d9d9"
#: past this many classes the table keeps the biggest and says so
_MAX_TABLE_ROWS = 26


def _seq_color(fraction: float) -> str:
    """A step of the sequential ramp for a 0..1 share."""
    if fraction <= 0:
        return "#ffffff"
    index = min(len(_SEQ) - 1, int(round(fraction * (len(_SEQ) - 1))))
    return _SEQ[index]


def _on_seq(fraction: float) -> str:
    """Ink that stays legible on that step."""
    return "#ffffff" if fraction >= 0.55 else _TEXT


def _draw_confusion(ax, analysis) -> None:
    """The matrix as a grid of shaded cells: rows are ground truth, columns
    predictions, the last of each is background (a miss / a false positive).

    Shaded by each cell's share of ITS ROW, so a rare class's error pattern is
    as readable as a common one's; the count is printed in every cell, so the
    exact number never depends on reading a colour."""
    names = list(analysis.classes)
    n = len(names)
    totals = [max(1, sum(row)) for row in analysis.matrix]
    ax.set_xlim(0, n)
    ax.set_ylim(0, n)
    ax.invert_yaxis()
    ax.set_aspect("equal")
    ax.axis("off")
    # the rotated column labels live above the grid; matplotlib does not
    # reserve room for text drawn outside the limits, so the caller leaves it
    ax.margins(0)
    size = max(5.0, min(9.0, 150.0 / n))

    for r, row in enumerate(analysis.matrix):
        for c, value in enumerate(row):
            share = value / totals[r]
            is_bg_cell = r == n - 1 and c == n - 1
            if is_bg_cell:
                continue  # background→background is not a thing
            ax.add_patch(_rect(c, r, _seq_color(share) if value else "#ffffff"))
            if r == c and r < n - 1:  # correct: outlined, never recoloured
                ax.add_patch(_rect(c, r, "none", edge=_DIAG, lw=1.6))
            if value:
                ax.text(c + 0.5, r + 0.5, str(value), ha="center", va="center",
                        fontsize=size, color=_on_seq(share))
    label = max(5.0, min(9.0, 140.0 / n))
    for i, name in enumerate(names):
        shown = name if len(name) <= 16 else name[:15] + "…"
        colour = _TEXT if i < n - 1 else _DIM
        ax.text(-0.3, i + 0.5, shown, ha="right", va="center", fontsize=label,
                color=colour)
        # upright, not slanted: 19 slanted labels run into each other, and a
        # vertical column of them costs a fixed, predictable strip of height.
        # Cut harder than the row labels: here length is height, and the strip
        # has to stay clear of the captions above it
        ax.text(i + 0.5, -0.3, name if len(name) <= 11 else name[:10] + "…",
                ha="center", va="bottom", fontsize=label, rotation=90, color=colour)
    ax.plot([0, n], [n - 1, n - 1], color=_RULE, lw=1)   # background row off
    ax.plot([n - 1, n - 1], [0, n], color=_RULE, lw=1)   # background column off


def _rect(col: int, row: int, face: str, *, edge: str = "none", lw: float = 0):
    from matplotlib.patches import Rectangle

    return Rectangle((col + 0.04, row + 0.04), 0.92, 0.92, facecolor=face,
                     edgecolor=edge, linewidth=lw)


def _draw_class_table(ax, analysis, ap_by_name: dict[str, float]) -> str:
    """Per-class performance, biggest class first. Recall and precision carry a
    light bar behind the number so the table can be scanned as well as read;
    the number is always there, so the bar adds nothing the text lacks.

    Returns a footnote when the table could not show every class."""
    rows = sorted(analysis.per_class, key=lambda c: (-c.instances, c.name))
    dropped = 0
    if len(rows) > _MAX_TABLE_ROWS:
        dropped = len(rows) - _MAX_TABLE_ROWS
        rows = rows[:_MAX_TABLE_ROWS]
    has_ap = bool(ap_by_name)
    columns = ["class", "boxes"] + (["AP@50"] if has_ap else []) + [
        "recall", "precision", "missed", "false", "wrong class",
    ]
    # x positions in axis space: the name column is wide, the rest are even
    xs = [0.0, 0.30] + ([0.40] if has_ap else [])
    start = xs[-1] + 0.10
    xs += [round(start + 0.115 * i, 3) for i in range(5)]

    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    height = 1.0 / (len(rows) + 1.6)
    size = max(6.0, min(9.5, 13.0 - 0.16 * len(rows)))
    top = 1.0

    for x, column in zip(xs, columns, strict=True):
        ax.text(x, top, column, fontsize=size - 0.5, color=_DIM,
                ha="left" if x == 0 else "right", va="top",
                transform=ax.transAxes)
    ax.plot([0, 1], [top - height * 0.55] * 2, color=_RULE, lw=0.8,
            transform=ax.transAxes)

    for i, cls in enumerate(rows):
        y = top - height * (i + 1.35)
        cells = [cls.name, str(cls.instances)]
        if has_ap:
            cells.append(f"{100 * ap_by_name.get(cls.name, 0.0):.0f}%")
        cells += [f"{100 * cls.recall:.0f}%", f"{100 * cls.precision:.0f}%",
                  str(cls.fn), str(cls.fp), str(cls.confused_as)]
        # the bars sit behind the two rate columns
        bar_at = 3 if has_ap else 2
        for offset, value in enumerate((cls.recall, cls.precision)):
            x = xs[bar_at + offset]
            ax.add_patch(_bar(ax, x - 0.105, y - height * 0.3, 0.105 * value, height * 0.62))
        for x, text in zip(xs, cells, strict=True):
            colour = _TEXT
            if text == cells[0] and len(text) > 20:
                text = text[:19] + "…"
            ax.text(x, y, text, fontsize=size, color=colour,
                    ha="left" if x == 0 else "right", va="center",
                    transform=ax.transAxes)
    return f"+ {dropped} more class(es) not shown" if dropped else ""


def _bar(ax, x: float, y: float, width: float, height: float):
    from matplotlib.patches import Rectangle

    return Rectangle((x, y), max(width, 0.0), height, facecolor=_SEQ[1],
                     edgecolor="none", transform=ax.transAxes, zorder=0)


def evaluation_figure(
    *,
    analysis,
    eval_report=None,
    advice=None,
    model: str = "",
):
    """One 16:9 sheet: the confusion matrix beside the per-class table (E6-T14).

    Both halves come from the SAME error analysis the evaluate page shows, at
    the threshold and IoU it was asked for, so the exported sheet and the page
    cannot disagree."""
    Figure, GridSpec = _import_matplotlib()
    fig = Figure(figsize=_FIGSIZE, dpi=_DPI)
    fig.patch.set_facecolor("white")
    # the matrix's column labels stand vertically above the grid, so the plot
    # row starts well below the header band rather than sharing its space
    gs = GridSpec(2, 2, figure=fig, height_ratios=[0.13, 1], width_ratios=[1, 1.1],
                  left=0.105, right=0.975, top=0.955, bottom=0.06,
                  hspace=0.62, wspace=0.09)

    head = fig.add_subplot(gs[0, :])
    head.axis("off")
    title = f"Evaluation — {analysis.split} split"
    head.text(0, 0.62, title, fontsize=21, fontweight="bold", color=_TEXT)
    facts = [f"run {analysis.run_id}"]
    if model:
        facts.append(model)
    facts.append(f"confidence ≥ {analysis.threshold:.2f}")
    facts.append(f"IoU ≥ {analysis.iou:.2f}")
    facts.append(f"{analysis.num_images} images")
    if eval_report is not None:
        facts.append(f"mAP@50 {100 * eval_report.map_50:.1f}%")
        facts.append(
            "current labels" if eval_report.labels == "current" else "run snapshot"
        )
    head.text(0, 0.12, "  ·  ".join(facts), fontsize=10.5, color=_DIM)
    totals = (f"{analysis.tp} correct   {analysis.fn} missed   "
              f"{analysis.fp} false   {analysis.confused} wrong class")
    head.text(1, 0.62, totals, fontsize=11.5, color=_TEXT, ha="right")
    if advice is not None and advice.confident:
        head.text(1, 0.12, f"suggested confidence {advice.recommended:.2f}",
                  fontsize=10.5, color=_DIM, ha="right")

    matrix_ax = fig.add_subplot(gs[1, 0])
    _draw_confusion(matrix_ax, analysis)
    table_ax = fig.add_subplot(gs[1, 1])
    ap_by_name = (
        {c.name: c.ap50 for c in eval_report.per_class} if eval_report is not None else {}
    )
    footnote = _draw_class_table(table_ax, analysis, ap_by_name)

    # section captions as figure text: an axes title would land on top of the
    # matrix's vertical column labels
    # below the header band, above the matrix's vertical column labels
    caption_y = 0.845
    fig.text(0.105, caption_y,
             "Confusion matrix — rows: ground truth, columns: predicted", fontsize=10,
             color=_TEXT)
    fig.text(0.105, caption_y - 0.025,
             "shaded by each cell's share of its row · green outline = correct",
             fontsize=9, color=_DIM)
    fig.text(0.545, caption_y, "Per class", fontsize=10, color=_TEXT)
    fig.text(0.545, caption_y - 0.025,
             "biggest class first · bars show recall and precision",
             fontsize=9, color=_DIM)

    tail = []
    if footnote:
        tail.append(footnote)
    if analysis.confused_pairs:
        worst = analysis.confused_pairs[0]
        tail.append(f"most confused: {worst.gt_name} → {worst.pred_name} ({worst.count})")
    fig.text(0.045, 0.028, "   ·   ".join(tail), fontsize=9, color=_DIM)
    fig.text(0.975, 0.028,
             f"horos {horos.__version__} · {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC}",
             fontsize=9, color=_DIM, ha="right")
    return fig
